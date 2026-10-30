"""Actual ES/MiniLM/upload routes. Run only with the owned run-upload-integration.py service."""
from dataclasses import dataclass, field, replace
from io import BytesIO
import hashlib
import json
import os
from pathlib import Path
import uuid

from docx import Document
from elasticsearch import Elasticsearch
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest


TEXT = "Controlled PDF conversion evidence fixture. Synthetic acceptance text only."


def pdf_bytes(text=TEXT):
    output = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=600)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 10 Tf 20 500 Td (' + text.encode('ascii') + b') Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(output)
    return output.getvalue()


@dataclass
class Actor:
    id: int
    headers: dict = field(repr=False)


@pytest.fixture(scope='module')
def isolated(tmp_path_factory):
    path = os.environ.get('SCHOLARLENS_ISOLATED_UPLOAD_CONFIG')
    if not path:
        pytest.skip('Use tools/audit/run-upload-integration.py; never point this test at an existing service')
    repo = Path(__file__).resolve().parents[2]
    config_path = Path(path).resolve(strict=True)
    assert config_path.is_relative_to((repo / 'tools/audit/reports').resolve(strict=True))
    config = json.loads(config_path.read_text(encoding='utf-8'))
    assert config['synthetic_data'] is True and config['output'] == str(config_path.parent)
    from backend.config import Settings
    settings = Settings.from_env()
    assert settings.url == config['url'] and settings.url.startswith('http://127.0.0.1:')
    assert settings.index == config['base_alias'] and settings.vector_index == config['vector_alias']
    assert str(config_path.parent / 'bootstrap-runtime') == settings.runtime_dir
    assert not any((settings.ai_api_key, settings.ai_model, settings.ai_base_url, settings.grobid_url, settings.pdf_root, settings.pdf_manifest_path))
    assert not os.environ.get('ELASTICSEARCH_API_KEY') and not os.environ.get('ELASTICSEARCH_CA_CERTS')
    es = Elasticsearch(settings.url, request_timeout=30, max_retries=0)
    info = es.info()
    assert info['cluster_name'] == config['cluster'] and info['cluster_uuid'] == config['cluster_uuid']
    assert info['version']['number'] == '8.19.23'
    # Import the default app only AFTER proving that its settings point to this new run.
    from backend.app1 import create_app
    from backend.import_jsonl import load_jsonl, import_records
    from backend.vector_index import build
    root = tmp_path_factory.mktemp('synthetic-base')
    settings = replace(settings, runtime_dir=str(root))
    source = root / 'explicitly-synthetic.jsonl'
    papers = [{'paper_id': 'fixture-acceptance-base-' + suffix,
               'title': 'Synthetic acceptance fixture ' + suffix,
               'fulltext': 'Explicit synthetic retrieval fixture. No research claims. Baseline invariant ' + suffix,
               'parse_status': 'partial'} for suffix in ('alpha', 'beta')]
    source.write_text(''.join(json.dumps(p) + '\n' for p in papers), encoding='utf-8')
    records, meta = load_jsonl(source, 'real')  # Production protocol mode, exclusively synthetic content.
    base = import_records(es, settings.index, records, meta)
    vectors = build(settings, root / 'vector-state.json', batch_size=4)
    evidence = {'synthetic_data': True, 'base_papers': len(papers), 'matrix': [], 'chains': [],
                'base_generation': base['generation'], 'base_vector_generation': vectors['generation']}
    try:
        yield {'es': es, 'settings': settings, 'create_app': create_app, 'papers': papers,
               'evidence': evidence, 'output': config_path.parent}
    finally:
        (config_path.parent / 'checks.json').write_text(json.dumps(evidence, indent=2), encoding='utf-8')
        es.close()


@pytest.fixture
def world(isolated, tmp_path):
    from backend.accounts import Store
    settings = isolated['settings']
    suffix = uuid.uuid4().hex[:8]
    settings = replace(settings, runtime_dir=str(tmp_path), uploads_index=settings.uploads_index.replace('-real', '-' + suffix + '-real'),
                       upload_vector_index=settings.upload_vector_index + '-' + suffix)
    store = Store(str(tmp_path))
    with TestClient(isolated['create_app'](settings, isolated['es'])) as api:
        actors = {}
        for role in ('owner', 'other', 'admin'):
            credentials = {'username': 'fixture_' + role, 'password': 'isolated-upload-password'}
            registered = api.post('/api/auth/register', json=credentials)
            assert registered.status_code == 201
            logged_in = api.post('/api/auth/login', json=credentials)
            assert logged_in.status_code == 200
            actors[role] = Actor(registered.json()['id'], {'Authorization': 'Bearer ' + logged_in.json()['access_token']})
        with store.connect() as db:
            db.execute("UPDATE users SET role='admin' WHERE id=?", (actors['admin'].id,))
        yield {**isolated, 'api': api, 'actors': actors, 'store': store, 'settings': settings}


def check(world, actor, method, path, expected, **kwargs):
    headers = dict(kwargs.pop('headers', {}))
    if actor:
        headers.update(world['actors'][actor].headers)
    response = world['api'].request(method, path, headers=headers, **kwargs)
    world['evidence']['matrix'].append({'actor': actor or 'guest', 'method': method, 'path': path,
                                       'expected': expected, 'actual': response.status_code})
    assert response.status_code == expected, {'actor': actor, 'method': method, 'path': path, 'actual': response.status_code}
    return response


def upload(world, actor='owner', content=None, filename='synthetic-fixture.pdf'):
    return check(world, actor, 'POST', '/api/uploads', 201,
                 files={'file': (filename, pdf_bytes() if content is None else content, 'application/pdf')}).json()


def test_real_upload_parse_vectors_retrieval_and_download_chain(world):
    es, settings = world['es'], world['settings']
    base_generation = list(es.indices.get_alias(name=settings.index))
    base_vectors = list(es.indices.get_alias(name=settings.vector_index))
    item = upload(world, filename='../../synthetic-fixture.pdf')
    private = '/api/uploads/' + item['id']
    route = '/api/admin/uploads/' + item['id']
    public_id = 'upload-sha256:' + hashlib.sha256(pdf_bytes()).hexdigest()
    check(world, None, 'GET', '/api/papers/' + public_id, 404)
    check(world, 'admin', 'POST', route + '/process', 409)
    check(world, 'admin', 'POST', route + '/review', 200, json={'decision': 'approve'})
    check(world, 'admin', 'POST', route + '/process', 202)
    result = check(world, 'owner', 'GET', private + '/result', 200).json()
    assert result['status'] == 'indexed' and result['paper_id'] == public_id, result
    paper = check(world, None, 'GET', '/api/papers/' + public_id, 200).json()
    assert TEXT in paper['fulltext'] and paper['authors'] == [] and paper['year'] is None and paper['date'] == ''
    assert paper['pdf_available'] is True
    for mode in ('keyword', 'semantic', 'hybrid'):
        query = check(world, None, 'GET', '/api/search', 200,
                      params={'q': 'Controlled PDF conversion evidence fixture', 'mode': mode, 'page_size': 10}).json()
        assert public_id in [r['paper_id'] for r in query['results']], {'mode': mode, 'results': query['results']}
    evidence = check(world, None, 'POST', '/api/evidence', 200,
                     json={'question': 'Controlled PDF conversion evidence fixture', 'paper_ids': [public_id], 'mode': 'hybrid', 'top_k': 5}).json()
    assert evidence and all(e['paper_id'] == public_id for e in evidence)
    chunk = es.search(index=settings.upload_vector_index, size=1)['hits']['hits'][0]['_source']
    assert len(chunk['embedding']) == 384
    fetched = check(world, None, 'GET', '/api/chunks/' + chunk['chunk_id'], 200).json()
    assert fetched['text'] == chunk['text'] and 'embedding' not in fetched
    original = check(world, None, 'GET', '/api/papers/' + public_id + '/pdf', 200)
    assert hashlib.sha256(original.content).hexdigest() == item['sha256']
    ranged = check(world, None, 'GET', '/api/papers/' + public_id + '/pdf', 206, headers={'Range': 'bytes=0-63'})
    assert ranged.content == original.content[:64] and ranged.headers['content-range'].startswith('bytes 0-63/')
    converted = check(world, None, 'GET', '/api/papers/' + public_id + '/docx', 200)
    assert TEXT in '\n'.join(p.text for p in Document(BytesIO(converted.content)).paragraphs)
    assert not list((Path(settings.runtime_dir) / 'docx').glob('job-*'))
    previous_upload_vector = list(es.indices.get_alias(name=settings.upload_vector_index))
    rebuilt = check(world, 'admin', 'POST', '/api/admin/uploads/vectorize', 200).json()
    assert rebuilt['reused_chunks'] > 0 and rebuilt['embedded_chunks'] == 0
    assert all(es.indices.exists(index=g) for g in previous_upload_vector)
    assert check(world, 'owner', 'GET', '/api/me/notifications', 200).json()['total'] >= 2
    archived = check(world, 'admin', 'DELETE', route, 200).json()
    assert archived['source_preserved'] and archived['indexed_document_preserved']
    check(world, 'admin', 'POST', route + '/process', 409)
    assert check(world, 'owner', 'GET', '/api/uploads', 200).json()['total'] == 0
    check(world, None, 'GET', '/api/papers/' + public_id + '/pdf', 200)
    assert es.count(index=settings.index)['count'] == 2
    assert list(es.indices.get_alias(name=settings.index)) == base_generation
    assert list(es.indices.get_alias(name=settings.vector_index)) == base_vectors
    world['evidence']['chains'].append({'name': 'actual_upload_to_vector_search_download_archive', 'passed': True,
        'synthetic_pdf': True, 'actual_embeddings': True, 'keyword_semantic_hybrid': True,
        'pdf_range': 206, 'docx_reopened': True, 'base_generation_preserved': True})


def test_private_upload_and_personal_record_ownership_matrix(world):
    item = upload(world)
    private = '/api/uploads/' + item['id']
    admin_route = '/api/admin/uploads/' + item['id']
    for suffix in ('', '/result', '/pdf', '/docx'):
        for actor, expected in ((None, 401), ('other', 404), ('owner', 200), ('admin', 200)):
            response = check(world, actor, 'GET', private + suffix, expected)
            if suffix == '/docx' and expected == 200:
                assert TEXT in '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    for actor, expected in ((None, 401), ('owner', 403), ('other', 403)):
        check(world, actor, 'GET', '/api/admin/uploads', expected)
        check(world, actor, 'GET', admin_route + '/pdf', expected)
        check(world, actor, 'POST', admin_route + '/review', expected, json={'decision': 'approve'})
        check(world, actor, 'POST', admin_route + '/process', expected)
        check(world, actor, 'POST', '/api/admin/uploads/vectorize', expected)
        check(world, actor, 'DELETE', admin_route, expected)
    assert check(world, 'other', 'GET', '/api/me/uploads', 200).json()['total'] == 0
    assert check(world, 'other', 'GET', '/api/me/papers', 200).json()['total'] == 0
    stored = list((world['store'].root / 'uploads/files').glob('*'))
    assert len(stored) == 1 and stored[0].stem == item['id']
    check(world, 'owner', 'POST', '/api/uploads', 409, files={'file': ('again.pdf', pdf_bytes(), 'application/pdf')})
    assert list((world['store'].root / 'uploads/files').glob('*')) == stored
    same_bytes = upload(world, actor='other')
    check(world, 'owner', 'GET', '/api/uploads/' + same_bytes['id'], 404)
    check(world, 'admin', 'POST', admin_route + '/review', 200, json={'decision': 'reject', 'reason': 'Synthetic fixture review'})
    check(world, 'admin', 'POST', admin_route + '/process', 409)
    check(world, 'admin', 'POST', admin_route + '/review', 409, json={'decision': 'approve'})
    notification = check(world, 'owner', 'GET', '/api/me/notifications', 200).json()['results'][0]
    for actor in ('other', 'admin'):
        assert check(world, actor, 'GET', '/api/me/notifications', 200).json()['total'] == 0
        check(world, actor, 'POST', f'/api/me/notifications/{notification["id"]}/read', 404)
    check(world, 'owner', 'POST', f'/api/me/notifications/{notification["id"]}/read', 200)
    paper_id = world['papers'][0]['paper_id']
    folder = check(world, 'owner', 'POST', '/api/me/favorite-folders', 201, json={'name': 'Private fixture folder'}).json()
    path = '/api/me/favorite-folders/' + str(folder['id'])
    check(world, 'owner', 'PUT', path + '/papers', 200, json={'paper_id': paper_id})
    for actor in ('other', 'admin'):
        assert check(world, actor, 'GET', '/api/me/favorite-folders', 200).json() == []
        check(world, actor, 'GET', path + '/papers', 404)
        check(world, actor, 'PATCH', path, 404, json={'name': 'Unauthorized rename'})
        check(world, actor, 'DELETE', path, 404)
        check(world, actor, 'PUT', path + '/papers', 404, json={'paper_id': paper_id})
        check(world, actor, 'DELETE', path + '/papers/' + paper_id, 404)
    assert check(world, 'owner', 'GET', path + '/papers', 200).json()['total'] == 1
    for actor in ('other', 'admin'):
        assert check(world, actor, 'GET', '/api/me/favorites', 200).json()['total'] == 0
        assert check(world, actor, 'GET', '/api/me/favorites/' + paper_id, 200).json()['is_favorite'] is False
        check(world, actor, 'DELETE', '/api/me/favorites/' + paper_id, 200)
        assert check(world, 'owner', 'GET', '/api/me/favorites/' + paper_id, 200).json()['is_favorite'] is True
    history = check(world, 'owner', 'POST', '/api/me/history', 201, json={'q': 'private fixture', 'field': 'all', 'mode': 'keyword'}).json()
    check(world, 'owner', 'POST', '/api/me/browsing-history', 201, json={'paper_id': paper_id})
    browsing = check(world, 'owner', 'GET', '/api/me/browsing-history', 200).json()['results'][0]
    for actor in ('other', 'admin'):
        for name, record in (('history', history), ('browsing-history', browsing)):
            assert check(world, actor, 'GET', '/api/me/' + name, 200).json()['total'] == 0
            check(world, actor, 'DELETE', f'/api/me/{name}/{record["id"]}', 404)
    for name, record in (('history', history), ('browsing-history', browsing)):
        assert check(world, 'owner', 'GET', '/api/me/' + name, 200).json()['total'] == 1
        check(world, 'owner', 'DELETE', f'/api/me/{name}/{record["id"]}', 200)
        assert check(world, 'owner', 'GET', '/api/me/' + name, 200).json()['total'] == 0
    workspace = check(world, 'owner', 'GET', '/api/me/workspace', 200).json()
    payload = {'version': workspace['version'], 'collections': ['全部收藏', 'Private fixture folder'],
               'favorites': [{'paper_id': paper_id, 'collection': 'Private fixture folder', 'note': 'Private fixture note'}], 'history': []}
    check(world, 'owner', 'PUT', '/api/me/workspace', 200, json=payload)
    for actor in ('other', 'admin'):
        assert check(world, actor, 'GET', '/api/me/workspace', 200).json()['favorites'] == []
    assert check(world, 'owner', 'GET', '/api/me/workspace', 200).json()['favorites'][0]['note'] == 'Private fixture note'


def test_conversation_owner_required_even_for_admin_and_all_generation_entrypoints(world):
    conversation = check(world, 'owner', 'POST', '/api/ai/conversations', 201, json={'title': 'Private fixture conversation'}).json()
    path = '/api/ai/conversations/' + conversation['id']
    paper_ids = [p['paper_id'] for p in world['papers']]
    body = {'question': 'Synthetic acceptance fixture', 'conversation_id': conversation['id'], 'mode': 'keyword', 'top_k': 5}
    for actor, expected in ((None, 401), ('other', 404), ('admin', 404)):
        check(world, actor, 'GET', path + '/messages', expected)
        check(world, actor, 'DELETE', path, expected)
        for endpoint, extra in (('/ask', {}), ('/ai/summarize', {'paper_ids': paper_ids[:1]}),
                                ('/ai/compare', {'paper_ids': paper_ids}), ('/agent/run', {'operation': 'evidence'})):
            check(world, actor, 'POST', '/api' + endpoint, expected, json={**body, **extra})
    for actor in ('other', 'admin'):
        assert check(world, actor, 'GET', '/api/ai/conversations', 200).json() == []
    check(world, 'owner', 'POST', '/api/ask', 503, json=body)
    assert check(world, 'owner', 'GET', path + '/messages', 200).json() == []
    result = check(world, 'owner', 'POST', '/api/agent/run', 200, json={**body, 'operation': 'evidence'}).json()
    assert result['evidence'] and result['model']['status'] == 'not_called'
    messages = check(world, 'owner', 'GET', path + '/messages', 200).json()
    assert len(messages) == 2 and messages[1]['citations']
    check(world, 'owner', 'DELETE', path, 200)
    check(world, 'owner', 'GET', path + '/messages', 404)
    with world['store'].connect() as db:
        assert db.execute('SELECT count(*) FROM conversation_messages WHERE conversation_id=?', (conversation['id'],)).fetchone()[0] == 0


def test_private_pdf_and_docx_responses_disable_caching(world):
    item = upload(world)
    for actor, path in (('owner', '/api/uploads/' + item['id'] + '/pdf'),
                        ('admin', '/api/admin/uploads/' + item['id'] + '/pdf'),
                        ('admin', '/api/uploads/' + item['id'] + '/docx')):
        response = check(world, actor, 'GET', path, 200)
        assert 'no-store' in response.headers.get('cache-control', ''), {'path': path, 'cache_control': response.headers.get('cache-control')}


def test_upload_rejections_scan_and_integrity_failure_preserve_sources_and_base(world):
    storage = world['store'].root / 'uploads/files'
    for actor, expected in ((None, 401), ('owner', 413)):
        check(world, actor, 'POST', '/api/uploads', expected, content=b'test', headers={'Content-Length': str(22 * 1024 * 1024)})
        assert list(storage.iterdir()) == []
    for name, content in (('fixture.txt', b'not a PDF'), ('fixture.pdf', b'not a PDF')):
        check(world, 'owner', 'POST', '/api/uploads', 400, files={'file': (name, content, 'application/pdf')})
        assert list(storage.iterdir()) == []
    # Exercise the file-size guard below the multipart request-size ceiling.
    oversized = pdf_bytes() + b' ' * (20 * 1024 * 1024 + 1 - len(pdf_bytes()))
    check(world, 'owner', 'POST', '/api/uploads', 413, files={'file': ('oversized.pdf', oversized, 'application/pdf')})
    assert list(storage.iterdir()) == []
    # A streamed body has no Content-Length; the receive guard must still enforce its bound.
    boundary = 'synthetic-acceptance-boundary'
    def streamed_body():
        yield ('--' + boundary + '\r\nContent-Disposition: form-data; name="file"; filename="stream.pdf"\r\nContent-Type: application/pdf\r\n\r\n').encode()
        for _ in range(22):
            yield b'x' * (1024 * 1024)
        yield ('\r\n--' + boundary + '--\r\n').encode()
    check(world, 'owner', 'POST', '/api/uploads', 413, content=streamed_body(), headers={'Content-Type': 'multipart/form-data; boundary=' + boundary})
    assert list(storage.iterdir()) == []
    scan = check(world, 'admin', 'POST', '/api/admin/uploads', 201,
                 files={'file': ('synthetic-scan.pdf', pdf_bytes(''), 'application/pdf')}).json()
    route = '/api/admin/uploads/' + scan['id']
    check(world, 'admin', 'POST', route + '/process', 202)
    failed = check(world, 'admin', 'GET', '/api/uploads/' + scan['id'], 200).json()
    assert failed['status'] == 'failed' and failed['reason'] == 'pdf_requires_ocr'
    check(world, 'admin', 'GET', route + '/pdf', 200)
    assert (storage / (scan['id'] + '.pdf')).is_file()
    damaged = upload(world, content=pdf_bytes('Synthetic integrity fixture text.'))
    source = storage / (damaged['id'] + '.pdf')
    source.write_bytes(source.read_bytes() + b'\nmodified fixture source\n')
    private = '/api/uploads/' + damaged['id']
    route = '/api/admin/uploads/' + damaged['id']
    check(world, 'owner', 'GET', private + '/pdf', 404)
    check(world, 'admin', 'GET', route + '/pdf', 409)
    check(world, 'owner', 'GET', private + '/docx', 409)
    check(world, 'admin', 'POST', route + '/review', 200, json={'decision': 'approve'})
    check(world, 'admin', 'POST', route + '/process', 202)
    failed = check(world, 'owner', 'GET', private, 200).json()
    assert failed['status'] == 'failed' and failed['reason'] == 'processing_failed'
    assert source.is_file()
    assert not world['es'].indices.exists_alias(name=world['settings'].uploads_index)
    assert not world['es'].indices.exists_alias(name=world['settings'].upload_vector_index)
    assert world['es'].count(index=world['settings'].index)['count'] == 2
    world['evidence']['chains'].append({'name': 'scan_and_hash_failure_preserve_sources_and_base', 'passed': True})
